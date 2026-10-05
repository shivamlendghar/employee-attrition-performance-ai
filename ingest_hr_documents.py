# ==========================================================
# ingest_hr_documents.py
#
# HR Document Ingestion + Text Extraction
#
# Supported:
#   PDF
#   DOCX
#   TXT
# ==========================================================

from pathlib import Path
import json
import re

from pypdf import PdfReader
from docx import Document


# ==========================================================
# PROJECT PATHS
# ==========================================================

PROJECT_DIR = Path(__file__).resolve().parent

HR_DOCUMENT_DIR = PROJECT_DIR / "hr_documents"

EXTRACTED_DIR = PROJECT_DIR / "extracted_documents"

DOCUMENT_INDEX_FILE = (
    EXTRACTED_DIR / "document_metadata.json"
)


# Create required directories
HR_DOCUMENT_DIR.mkdir(exist_ok=True)
EXTRACTED_DIR.mkdir(exist_ok=True)


# ==========================================================
# SUPPORTED FILE TYPES
# ==========================================================

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
}


# ==========================================================
# TEXT CLEANING
# ==========================================================

def clean_text(text):
    """
    Clean extracted document text.

    Operations:
        - Replace tabs/newlines with spaces
        - Remove repeated whitespace
        - Remove strange control characters
        - Strip leading/trailing spaces
    """

    if not text:
        return ""

    # Convert tabs/newlines to spaces
    text = text.replace("\n", " ")
    text = text.replace("\r", " ")
    text = text.replace("\t", " ")

    # Remove control characters
    text = re.sub(
        r"[\x00-\x08\x0B\x0C\x0E-\x1F\x7F]",
        " ",
        text
    )

    # Remove repeated spaces
    text = re.sub(
        r"\s+",
        " ",
        text
    )

    return text.strip()


# ==========================================================
# PDF EXTRACTION
# ==========================================================

def extract_pdf(file_path):
    """
    Extract text from a PDF file.
    """

    reader = PdfReader(
        str(file_path)
    )

    pages = []

    for page_number, page in enumerate(
        reader.pages,
        start=1
    ):

        try:

            text = page.extract_text()

        except Exception as error:

            print(
                f"Warning: Could not extract "
                f"page {page_number} from "
                f"{file_path.name}: {error}"
            )

            text = ""

        if text:

            pages.append(
                f"Page {page_number}: {text}"
            )

    return "\n".join(pages)


# ==========================================================
# DOCX EXTRACTION
# ==========================================================

def extract_docx(file_path):
    """
    Extract paragraphs and table text from DOCX.
    """

    document = Document(
        str(file_path)
    )

    content = []

    # Extract paragraphs
    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:

            content.append(text)

    # Extract tables
    for table in document.tables:

        for row in table.rows:

            row_text = []

            for cell in row.cells:

                cell_text = cell.text.strip()

                if cell_text:

                    row_text.append(cell_text)

            if row_text:

                content.append(
                    " | ".join(row_text)
                )

    return "\n".join(content)


# ==========================================================
# TXT EXTRACTION
# ==========================================================

def extract_txt(file_path):
    """
    Read a TXT file.
    """

    # utf-8-sig handles normal UTF-8 as well as files
    # containing a UTF-8 BOM.
    with open(
        file_path,
        "r",
        encoding="utf-8-sig",
        errors="replace"
    ) as file:

        return file.read()


# ==========================================================
# GENERAL EXTRACTION
# ==========================================================

def extract_text(file_path):
    """
    Extract text based on file extension.
    """

    extension = (
        file_path.suffix.lower()
    )

    if extension == ".pdf":

        return extract_pdf(
            file_path
        )

    elif extension == ".docx":

        return extract_docx(
            file_path
        )

    elif extension == ".txt":

        return extract_txt(
            file_path
        )

    else:

        raise ValueError(
            f"Unsupported file type: "
            f"{extension}"
        )


# ==========================================================
# CATEGORY DETECTION
# ==========================================================

def determine_category(file_path):
    """
    Determine document category based on
    its parent folder.

    Example:

        hr_documents/leave/leave_policy.pdf

    becomes:

        leave
    """

    relative_path = file_path.relative_to(
        HR_DOCUMENT_DIR
    )

    parts = relative_path.parts

    if len(parts) >= 2:

        return parts[0].lower()

    return "general"


# ==========================================================
# SAFE FILE NAME
# ==========================================================

def create_safe_filename(file_path):
    """
    Convert original filename to a safe output filename.
    """

    name = file_path.stem

    name = re.sub(
        r"[^a-zA-Z0-9_-]+",
        "_",
        name
    )

    return name.lower()


# ==========================================================
# PROCESS ONE DOCUMENT
# ==========================================================

def process_document(file_path):
    """
    Extract and clean one document.
    """

    print(
        f"\nProcessing: {file_path}"
    )

    try:

        raw_text = extract_text(
            file_path
        )

        cleaned_text = clean_text(
            raw_text
        )

        if not cleaned_text:

            print(
                "  WARNING: No text extracted."
            )

            return None

        category = determine_category(
            file_path
        )

        safe_name = create_safe_filename(
            file_path
        )

        output_file = (
            EXTRACTED_DIR
            / f"{safe_name}.txt"
        )

        # Save extracted text
        with open(
            output_file,
            "w",
            encoding="utf-8"
        ) as file:

            file.write(
                cleaned_text
            )

        metadata = {
            "document_id": safe_name,
            "original_filename": file_path.name,
            "source_path": str(
                file_path.relative_to(
                    PROJECT_DIR
                )
            ),
            "file_type": file_path.suffix.lower(),
            "category": category,
            "character_count": len(
                cleaned_text
            ),
            "word_count": len(
                cleaned_text.split()
            ),
            "extracted_text_file": str(
                output_file.relative_to(
                    PROJECT_DIR
                )
            ),
        }

        print(
            f"  Category       : {category}"
        )

        print(
            f"  Words          : "
            f"{metadata['word_count']}"
        )

        print(
            f"  Characters     : "
            f"{metadata['character_count']}"
        )

        print(
            f"  Saved          : {output_file}"
        )

        return metadata

    except Exception as error:

        print(
            f"  ERROR: {error}"
        )

        return None


# ==========================================================
# FIND DOCUMENTS
# ==========================================================

def find_documents():
    """
    Find all supported HR documents recursively.
    """

    documents = []

    for file_path in HR_DOCUMENT_DIR.rglob("*"):

        if not file_path.is_file():
            continue

        if (
            file_path.suffix.lower()
            in SUPPORTED_EXTENSIONS
        ):

            documents.append(
                file_path
            )

    return sorted(
        documents
    )


# ==========================================================
# SAVE METADATA
# ==========================================================

def save_metadata(metadata):
    """
    Save document metadata as JSON.
    """

    with open(
        DOCUMENT_INDEX_FILE,
        "w",
        encoding="utf-8"
    ) as file:

        json.dump(
            metadata,
            file,
            indent=4,
            ensure_ascii=False
        )


# ==========================================================
# MAIN
# ==========================================================

def main():

    print("=" * 70)
    print("HR DOCUMENT INGESTION SYSTEM")
    print("=" * 70)

    documents = find_documents()

    print(
        f"\nDocuments found: {len(documents)}"
    )

    if not documents:

        print("\nNo HR documents found.")

        print(
            "\nAdd PDF, DOCX or TXT files inside:"
        )

        print(
            HR_DOCUMENT_DIR
        )

        print(
            "\nExample:"
        )

        print(
            "hr_documents/"
        )

        print(
            "    leave/"
        )

        print(
            "        leave_policy.pdf"
        )

        return

    metadata = []

    for document in documents:

        result = process_document(
            document
        )

        if result:

            metadata.append(
                result
            )

    save_metadata(
        metadata
    )

    # ======================================================
    # SUMMARY
    # ======================================================

    print("\n" + "=" * 70)
    print("INGESTION COMPLETED")
    print("=" * 70)

    print(
        f"Successfully processed: "
        f"{len(metadata)}"
    )

    print(
        f"Failed/skipped: "
        f"{len(documents) - len(metadata)}"
    )

    print(
        "\nExtracted documents:"
    )

    print(
        EXTRACTED_DIR
    )

    print(
        "\nMetadata:"
    )

    print(
        DOCUMENT_INDEX_FILE
    )


# ==========================================================
# ENTRY POINT
# ==========================================================

if __name__ == "__main__":
    main()